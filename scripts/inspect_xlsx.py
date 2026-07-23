"""Inspect the structure of db_surdi.xlsx"""
import openpyxl
from pathlib import Path

path = Path("/home/z/my-project/upload/db_surdi.xlsx")
wb = openpyxl.load_workbook(path, data_only=True)

print("=" * 80)
print(f"Sheet names: {wb.sheetnames}")
print("=" * 80)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n### Sheet: {sheet_name} ###")
    print(f"  Dimensions: {ws.dimensions}")
    print(f"  Max row: {ws.max_row}, Max col: {ws.max_column}")

    # Print first 8 rows
    print("\n  First 8 rows:")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True)):
        print(f"    Row {i+1}: {row}")

    # Print header row in detail
    print("\n  Header columns:")
    header = [cell.value for cell in ws[1]]
    for idx, h in enumerate(header, 1):
        print(f"    Col {idx}: {h}")
    print("-" * 80)
